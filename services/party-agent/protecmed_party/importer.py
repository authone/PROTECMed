"""Local XLSX / strict canonical CSV import for the party agent (blueprint 3.1-3.3, 3.6).

Positive allowlist: only the five mapped columns of the one approved sheet are read
into the canonical table. No full patient frame is ever materialized, no identifier,
date, note or free-text column is copied, and no formula is evaluated. External-link
parts are counted and never followed. Macros and encrypted workbooks are rejected.

Two explicit formula modes, per blueprint 3.2:
  literal-only            default; a formula in a selected clinical cell is rejected.
  reviewed-cached-values  formula cells allowed only with a local operator
                          acknowledgement bound to this exact source digest; every
                          selected formula cell must carry a non-error, nonempty
                          saved result. Cache freshness is NOT established by this.
"""
from __future__ import annotations
import csv
import hashlib
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterator
from zipfile import BadZipFile, ZipFile

from .canonical import FIELDS, header_key, normalize, project_raw_rows
from .errors import ImportRejected

LITERAL_ONLY = "literal-only"
REVIEWED_CACHED = "reviewed-cached-values"
MODES = (LITERAL_ONLY, REVIEWED_CACHED)

MAX_SOURCE_BYTES = 64 * 1024**2
MAX_EXPANDED_BYTES = 128 * 1024**2
MAX_SHEET_ROWS = 10002
EXCEL_ERRORS = {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!",
                "#N/A", "#GETTING_DATA", "#SPILL!", "#CALC!"}


@dataclass(frozen=True)
class Mapping:
    mapping_id: str
    sheet: str
    headers: dict[str, str]
    local_row_cap: int
    digest: str


def load_mapping(path: Path) -> Mapping:
    raw = Path(path).read_bytes()
    document = json.loads(raw.decode("utf-8"))
    if type(document) is not dict or document.get("schema_version") != "2.0":
        raise ImportRejected("MAPPING_VERSION")
    fields = document.get("fields")
    if type(fields) is not dict or set(fields) != set(FIELDS):
        raise ImportRejected("MAPPING_FIELDS")
    if any(type(v) is not str or not v.strip() for v in fields.values()):
        raise ImportRejected("MAPPING_HEADERS")
    if document.get("external_link_refresh") is not False:
        raise ImportRejected("EXTERNAL_LINK_REFRESH_FORBIDDEN")
    sheet = document.get("sheet")
    if type(sheet) is not str or not sheet:
        raise ImportRejected("MAPPING_SHEET")
    cap = document.get("local_row_cap")
    if type(cap) is not int or not 0 < cap <= 10000:
        raise ImportRejected("MAPPING_ROW_CAP")
    # A mapping version change changes this digest and therefore requires a new run.
    digest = hashlib.sha256(json.dumps(document, sort_keys=True, separators=(",", ":"),
                                       ensure_ascii=False).encode("utf-8")).hexdigest()
    return Mapping(document["mapping_id"], sheet, dict(fields), cap, digest)


def file_digest(path: Path) -> str:
    """Local-only SHA-256 of the clinical source. Never leaves the provider."""
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _check_package(path: Path) -> dict[str, int]:
    """Structural guards on the real container, not on the filename extension."""
    if Path(path).stat().st_size > MAX_SOURCE_BYTES:
        raise ImportRejected("XLSX_SIZE")
    try:
        with ZipFile(path) as archive:
            infos = archive.infolist()
            names = [i.filename for i in infos]
            if len(names) != len(set(names)):
                raise ImportRejected("DUPLICATE_ZIP_ENTRY")
            if sum(i.file_size for i in infos) > MAX_EXPANDED_BYTES:
                raise ImportRejected("XLSX_EXPANDED_SIZE")
            lowered = [n.lower() for n in names]
            if any("vbaproject" in n for n in lowered):
                raise ImportRejected("MACRO_WORKBOOK")
            if any(n.startswith("encryptedpackage") for n in lowered):
                raise ImportRejected("ENCRYPTED_WORKBOOK")
            if not any(n == "xl/workbook.xml" for n in lowered):
                raise ImportRejected("NOT_AN_XLSX_PACKAGE")
            external = sum(1 for n in lowered
                           if n.startswith("xl/externallinks/") and n.endswith(".xml"))
    except BadZipFile:
        # An OOXML-encrypted workbook is an OLE container, not a ZIP.
        raise ImportRejected("ENCRYPTED_OR_INVALID_WORKBOOK") from None
    return {"external_link_parts_not_refreshed": external}


def _is_formula(value: Any, data_type: Any) -> bool:
    if data_type == "f":
        return True
    if type(value) is str and value.startswith("="):
        return True
    return type(value).__name__ in {"ArrayFormula", "DataTableFormula"}


def _cached_value(value: Any) -> Any:
    if type(value) is str and value.strip().upper() in EXCEL_ERRORS:
        raise ImportRejected("EXCEL_ERROR_CELL")
    return value


def _open_views(path: Path):
    import openpyxl  # imported lazily so the module loads without the optional dependency
    common = {"read_only": True, "data_only": None, "keep_vba": False, "keep_links": False}
    formulas = openpyxl.load_workbook(path, **{**common, "data_only": False})
    values = openpyxl.load_workbook(path, **{**common, "data_only": True})
    return formulas, values


def _select_sheet(book: Any, sheet: str) -> Any:
    matches = [name for name in book.sheetnames if header_key(name) == header_key(sheet)]
    if len(matches) != 1:
        raise ImportRejected("SHEET_MISSING_OR_AMBIGUOUS")
    return book[matches[0]]


def _column_map(header_cells: list[Any], mapping: Mapping) -> dict[str, int]:
    seen: dict[str, list[int]] = {}
    for index, cell in enumerate(header_cells):
        key = header_key(cell.value)
        if key:
            seen.setdefault(key, []).append(index)
    columns: dict[str, int] = {}
    for field, header in mapping.headers.items():
        found = seen.get(header_key(header), [])
        # Never fall back to the closest similarly named column.
        if len(found) != 1:
            raise ImportRejected("REQUIRED_HEADER_MISSING_OR_AMBIGUOUS",
                                 {"field": field, "matches": len(found)})
        columns[field] = found[0]
    return columns


def _pick(cells: list[Any], index: int) -> Any:
    return cells[index] if index < len(cells) else None


def import_workbook(path: Path, mapping: Mapping, mode: str = LITERAL_ONLY, *,
                    cache_acknowledgement: dict[str, Any] | None = None
                    ) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Return (canonical records, local import report). Fails closed on any violation."""
    if mode not in MODES:
        raise ImportRejected("IMPORT_MODE")
    path = Path(path)
    package = _check_package(path)
    source_digest = file_digest(path)
    if mode == REVIEWED_CACHED:
        _check_acknowledgement(cache_acknowledgement, source_digest, mapping)
    formulas_book, values_book = _open_views(path)
    try:
        formula_sheet = _select_sheet(formulas_book, mapping.sheet)
        value_sheet = _select_sheet(values_book, mapping.sheet)
        raw_rows, formula_counts = _read_rows(formula_sheet, value_sheet, mapping, mode)
    finally:
        formulas_book.close()
        values_book.close()
    records, report = project_raw_rows(raw_rows)
    if len(records) > mapping.local_row_cap:
        raise ImportRejected("LOCAL_ROW_LIMIT")
    report.update({
        "source": "xlsx",
        "sheet": mapping.sheet,
        "mapping_id": mapping.mapping_id,
        "mapping_digest": mapping.digest,
        "import_mode": mode,
        "selected_formula_cells": formula_counts,
        "cache_freshness_verified": False,
        **package,
    })
    if mode == REVIEWED_CACHED:
        report["cache_acknowledgement"] = dict(cache_acknowledgement or {})
        report["cache_warning"] = ("Saved formula results were accepted on operator "
                                   "acknowledgement; freshness cannot be established "
                                   "from the workbook alone.")
    return records, report


def _check_acknowledgement(acknowledgement: Any, source_digest: str, mapping: Mapping) -> None:
    if type(acknowledgement) is not dict:
        raise ImportRejected("CACHE_ACKNOWLEDGEMENT_REQUIRED")
    required = {"operator", "acknowledged_utc", "source_sha256", "mapping_id", "statement"}
    if not required <= set(acknowledgement):
        raise ImportRejected("CACHE_ACKNOWLEDGEMENT_FIELDS")
    if acknowledgement["source_sha256"] != source_digest:
        # The reviewed workbook is not this workbook: a changed file needs a new review.
        raise ImportRejected("CACHE_ACKNOWLEDGEMENT_SOURCE_MISMATCH")
    if acknowledgement["mapping_id"] != mapping.mapping_id:
        raise ImportRejected("CACHE_ACKNOWLEDGEMENT_MAPPING_MISMATCH")


def _read_rows(formula_sheet: Any, value_sheet: Any, mapping: Mapping,
               mode: str) -> tuple[list[dict[str, Any]], dict[str, int]]:
    formula_rows: Iterator[Any] = formula_sheet.iter_rows()
    value_rows: Iterator[Any] = value_sheet.iter_rows()
    try:
        formula_header = list(next(formula_rows))
        value_header = list(next(value_rows))
    except StopIteration:
        raise ImportRejected("HEADER_ROW") from None
    columns = _column_map(value_header, mapping)
    if _column_map(formula_header, mapping) != columns:
        raise ImportRejected("HEADER_VIEW_MISMATCH")
    formula_counts = {field: 0 for field in FIELDS}
    raw_rows: list[dict[str, Any]] = []
    for position, (formula_cells, value_cells) in enumerate(zip(formula_rows, value_rows), start=2):
        if position > MAX_SHEET_ROWS:
            raise ImportRejected("ROW_LIMIT")
        formula_cells, value_cells = list(formula_cells), list(value_cells)
        raw: dict[str, Any] = {}
        for field, index in columns.items():
            source = _pick(formula_cells, index)
            cached = _pick(value_cells, index)
            is_formula = source is not None and _is_formula(
                source.value, getattr(source, "data_type", None))
            if is_formula:
                formula_counts[field] += 1
                if mode == LITERAL_ONLY:
                    raise ImportRejected("FORMULA_NOT_ALLOWED", {"field": field})
            value = _cached_value(None if cached is None else cached.value)
            if is_formula and (value is None or (type(value) is str and value == "")):
                # A missing cache is never null, false, zero or a reason to drop the row.
                raise ImportRejected("FORMULA_CACHE_MISSING", {"field": field})
            raw[field] = value
        raw_rows.append(raw)
    return raw_rows, formula_counts


def read_canonical_csv(path: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Strict canonical CSV: exact five-column header, exact width, no extra columns.

    Every CSV record is already admitted, so an all-missing canonical row is kept.
    """
    records: list[dict[str, Any]] = []
    with Path(path).open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.reader(stream)
        header = next(reader, None)
        if header != list(FIELDS):
            raise ImportRejected("CANONICAL_HEADER_MISMATCH")
        for cells in reader:
            if len(cells) != len(FIELDS):
                raise ImportRejected("CANONICAL_ROW_WIDTH")
            records.append({f: normalize(f, v) for f, v in zip(FIELDS, cells)})
            if len(records) > 10000:
                raise ImportRejected("LOCAL_ROW_LIMIT")
    report = {"source": "canonical-csv", "admitted_rows": len(records),
              "skipped_blank_rows": 0, "import_mode": "canonical-csv",
              "selected_formula_cells": {f: 0 for f in FIELDS},
              "cache_freshness_verified": False}
    return records, report
